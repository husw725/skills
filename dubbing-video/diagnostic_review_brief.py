"""Build and publish the finite three-episode diagnostic review brief on Windows."""
import argparse
import hashlib
from pathlib import Path
import sys
from dub_episode import digest, safe_submit
from mflix_client import Mflix
from process_lock import process_lock
from translate import read_json, save_json, fingerprint
from cartesia_pilot import bind

NAME = 'Carmilla | EP02 v3 + EP07 v2 + EP29 v2 | Guia de revisão pt-BR'
ROOT = Path('output/review-publication/148/Carmilla/diagnostic-brief-20261016')


def stamp(seconds):
    return f'{int(seconds)//60:02}:{seconds%60:06.3f}'


def build():
    from docx import Document
    from docx.shared import Inches, Pt
    ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.65)
    normal = doc.styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10)
    doc.add_heading('Carmilla — revisão diagnóstica de dublagem pt-BR', 0)
    doc.add_paragraph('版本：EP02 v3、EP07 v2、EP29 v2。请完整听看三集，再按时间码反馈；已知疑点并非全部问题。')
    doc.add_paragraph('Objetivo: identificar problemas reais de português brasileiro, atuação, voz e sincronismo. '
                      'Assista aos episódios completos, inclusive trechos não destacados. '
                      'Estes são candidatos para revisão; não foram aprovados por escuta nativa.')
    doc.add_heading('Como registrar o feedback', 1)
    doc.add_paragraph('Informe: episódio + versão + tempo inicial/final + frase exata + categoria + '
                      'o que soa errado + sugestão (texto e/ou indicação de pronúncia/atuação). '
                      'Se possível, anexe uma referência falada curta. Registre também o que passou. '
                      'Sem comentário não significa aprovado. Use a versão indicada, evitando misturar comentários antigos.')
    for text in [
        'Tradução/contexto: sentido completo, negação, prazo, nomes, relações, registro e continuidade entre falas.',
        'Português brasileiro: sotaque/região, vogais nasais, tonicidade, consoantes, nomes e resíduos de inglês. '
        'Não alterar uma palavra corretamente escrita só para compensar um erro de pronúncia do sintetizador.',
        'Atuação: medo, grito, ameaça, ironia, intenção, respiração e intensidade coerentes com a cena.',
        'Identidade vocal: consistência do mesmo personagem entre frases/modelos; personagens distinguíveis.',
        'Tempo e mixagem: velocidade natural, finais completos, pausas, sincronia, sobreposição, '
        'nível do diálogo, música/efeitos, ruído ou inglês residual. Indicar a origem apenas se identificável.'
    ]: doc.add_paragraph(text, style='List Bullet')
    doc.add_paragraph('Formato sugerido: EP07 v2 | 00:09–00:11 | “CAMILA!” | atuação/pronúncia | '
                      '[problema observado] | [sugestão]. Separe erro textual de erro de áudio.')
    doc.add_heading('Método e limites do experimento', 1)
    doc.add_paragraph('Modelo híbrido: Cartesia nas falas principais; MiniMax em trechos selecionados. '
                      'Todas as vozes foram reutilizadas, sem nova clonagem. Texto e tempos das legendas originais '
                      'desta tradução foram preservados. Algumas frases consecutivas do mesmo personagem foram agrupadas. '
                      'Há aproveitamento de intervalo após certas legendas: avalie o ritmo no vídeo, não só o texto. '
                      'Parâmetro pt-BR/Portuguese, emoção solicitada e tarefa concluída não comprovam sotaque ou atuação corretos. '
                      'Não é um teste controlado para eleger um fornecedor: cada trecho tem apenas uma versão nova.')
    matrix = {'project_id': 148, 'date': '2026-10-16', 'native_review_passed': False,
              'comment_schema': ['episode', 'version', 'asset_id', 'comment_id', 'time_start', 'time_end',
                                 'quoted_phrase', 'category', 'observation', 'suggestion', 'verified_or_hypothesis'],
              'learning_policy': 'Anchor feedback to exact version/time/phrase; test existing voices on small samples; '
                                 'separate confirmed findings from hypotheses; no automatic cloning or season regeneration.',
              'episodes': []}
    cfg = read_json('output/diagnostic-review-config.json')
    for n, version, subdir, receipt in [(2, 3, 'feedback-v3', 'feedback-ep02-v3'),
                                      (7, 2, 'diagnostic-v2', 'diagnostic-ep07-v2'),
                                      (29, 2, 'diagnostic-v2', 'diagnostic-ep29-v2')]:
        ep_root = Path(f'output/drama-01/episode-{n:02}')
        report = read_json(ep_root/subdir/'render-report.json')
        publication = read_json(Path('output/review-publication/148/Carmilla')/receipt/'publication-result.json')
        if (publication['project_id'] != 148 or report['episode'] != n or publication['version'] != version
                or publication['sha256'] != report['output_sha256']):
            raise ValueError('Candidate publication differs from report')
        plan = read_json(ep_root/'dub-plan-v1.json'); source = {u['id']: u for u in plan['units']}
        covered = [c for u in report['utterances'] for c in u['cue_ids']]
        expected = [c for u in plan['units'] for c in u['cue_ids']]
        if len(covered) != len(expected) or len(set(covered)) != len(covered) or set(covered) != set(expected):
            raise ValueError('Bilingual reference coverage differs')
        doc.add_page_break(); doc.add_heading(f'EP{n:02} — versão {version} — Review {publication["asset_id"]}', 1)
        doc.add_paragraph('Vídeo: '+publication['url'])
        doc.add_paragraph('SHA256: '+report['output_sha256'])
        if n == 2:
            doc.add_paragraph('Somente os dois “não” de Laura foram substituídos por MiniMax (fearful; '
                              'controle fonético no texto de síntese; gasps no primeiro). '
                              'As outras 15 frases Cartesia v2 foram mantidas. '
                              'Verificar pronúncia nasal, medo/grito, timbre ao trocar de modelo e duração. '
                              'Os dois comentários anteriores não tinham tempo; estes são locais candidatos. '
                              '“Quinze anos” também precisa ser ouvido: ausência de nova crítica não é aprovação.')
            focus = [{'time': '00:49.4', 'check': 'Não... — desespero, nasalização, respiração e continuidade de Laura.'},
                     {'time': '00:55.0–00:56.7', 'check': 'NÃO! — grito de medo; áudio pode ultrapassar a janela da legenda.'}]
        else:
            focus = cfg['episodes'][str(n)]['review_focus']
        for f in focus: doc.add_paragraph(f['time']+' — '+f['check'], style='List Bullet')
        if n == 29:
            doc.add_paragraph('Atenção a 00:47: “How many years.” / “Quantos anos?” não corresponde ao roteiro '
                              '(“Her minions…”). A fala antiga foi mantida e não foi sintetizada novamente. '
                              'O falante na evidência continua desconhecido; a voz de produção provisória é Carmilla. '
                              'Confirmar áudio original, transcrição e personagem antes de corrigir. '
                              'Este vídeo corresponde ao EP-30 no storyboard final de 33 episódios; não confundir a numeração.')
        doc.add_heading('Referência bilíngue por grupo de dublagem', 2)
        doc.add_paragraph('Tempos abaixo são as janelas originais. A duração audível pode usar um intervalo posterior. '
                          'O texto à direita é a tradução, não os controles fonéticos usados na síntese.')
        table = doc.add_table(rows=1, cols=4); table.style = 'Table Grid'
        for c, text in zip(table.rows[0].cells, ['Tempo / personagem', 'Original (EN)', 'Tradução (pt-BR)', 'Áudio']): c.text = text
        entries = []
        for u in report['utterances']:
            ids = u.get('source_unit_ids') or [u['id']]
            english = ' '.join(source[k]['text'] for k in ids)
            provider = u.get('generation_provider', 'not-recorded')
            row = table.add_row().cells
            values = [stamp(u['start'])+'–'+stamp(u['end'])+'\n'+u['production_voice'],
                      english, u['translation'], provider]
            for cell, text in zip(row, values): cell.text = text
            entries.append({k: u.get(k) for k in ('id', 'cue_ids', 'source_unit_ids', 'start', 'end',
                            'production_voice', 'voice_id', 'generation_provider', 'task_id', 'synthesis_request',
                            'atempo_factor', 'audio_end_limit', 'post_subtitle_gap_used_seconds',
                            'emotion_control_applied', 'source_reader_verified')})
        matrix['episodes'].append({'episode': n, 'version': version, 'asset_id': publication['asset_id'],
                'output_sha256': report['output_sha256'], 'input_fingerprint': report['input_fingerprint'],
                'review_focus': focus, 'utterances': entries})
    doc.add_page_break(); doc.add_heading('Como usaremos os comentários', 1)
    doc.add_paragraph('Cada comentário será associado ao episódio, versão, frase e, quando fornecido, tempo. '
                      'Comentários sem tempo terão apenas candidatos de localização. '
                      'Vamos distinguir ajustes de tradução, pronúncia, atuação, tempo e mixagem; '
                      'registrar a causa como hipótese até haver evidência; e validar pequenos trechos antes de expandir. '
                      'Não haverá clonagem repetida nem regeneração de toda a série só por hipótese.')
    path = ROOT/'Carmilla_pt-BR_diagnostic_review_20261016.docx'
    doc.save(path); save_json(ROOT/'feedback-learning-matrix.private.json', matrix)
    return path


def publish(config):
    ROOT.mkdir(parents=True, exist_ok=True)
    with process_lock(ROOT/'publish.lock', timeout=0):
        path = ROOT/'Carmilla_pt-BR_diagnostic_review_20261016.docx'
        inputs = {'config': digest(Path('output/diagnostic-review-config.json'))}
        for n, subdir, receipt in [(2, 'feedback-v3', 'feedback-ep02-v3'),
                                  (7, 'diagnostic-v2', 'diagnostic-ep07-v2'),
                                  (29, 'diagnostic-v2', 'diagnostic-ep29-v2')]:
            for p in [Path(f'output/drama-01/episode-{n:02}')/subdir/'render-report.json',
                      Path(f'output/drama-01/episode-{n:02}/dub-plan-v1.json'),
                      Path('output/review-publication/148/Carmilla')/receipt/'publication-result.json']:
                inputs[str(p)] = digest(p)
        marker = ROOT/'input-fingerprint.json'
        if path.exists() and not marker.exists(): raise ValueError('Unbound existing brief')
        bind(marker, inputs)
        if not path.exists(): path = build()
        sha = digest(path); client = Mflix(config)
        def documents():
            data = client.call('getProjectReviewMaterials', {'projectId': 148, 'isNeedArtDirectionSubMaterials': False})
            if data['projectId'] != 148: raise ValueError('Unexpected project')
            return data['documentMaterials']
        before = documents(); save_json(ROOT/'documents-before.private.json', before)
        upload = ROOT/'upload.private.json'
        if upload.exists():
            uploaded = read_json(upload)
            if uploaded['sha256'] != sha: raise ValueError('Brief changed after upload')
        else:
            if list(ROOT.glob('create-submission*')): raise ValueError('Submission exists without upload receipt')
            url = client.upload(path)
            response = client.http.get(url, timeout=60)
            if response.status_code != 200 or hashlib.sha256(response.content).hexdigest() != sha:
                raise ValueError('Brief upload SHA readback failed')
            uploaded = {'url': url, 'sha256': sha}; save_json(upload, uploaded)
        same_name = [r for r in before if r['name'] == NAME]
        if same_name and (len(same_name) != 1 or same_name[0]['url'] != uploaded['url']):
            raise ValueError('Another brief with same name exists')
        if not same_name:
            material = {k: None for k in ('artDirectionSubMaterials', 'assetId', 'category', 'createdAt',
                        'durationSeconds', 'parentAssetId', 'status', 'tags', 'updatedAt', 'version')}
            material.update({'projectId': 148, 'gateType': 'DOCUMENT', 'documentKind': 'DOCUMENT',
                             'name': NAME, 'url': uploaded['url'],
                             'note': 'Carmilla; EP02 v3 + EP07 v2 + EP29 v2 diagnostic reviewer guide; '
                                     'bilingual EN/pt-BR reference, timecoded focus, categories and feedback template. '
                                     'Not native-approved. SHA256: '+sha})
            sig = ROOT/'request-fingerprint.json'; value = fingerprint(material)
            if sig.exists() and read_json(sig)['fingerprint'] != value: raise ValueError('Brief request changed')
            save_json(sig, {'fingerprint': value})
            safe_submit(client, 'createReviewMaterial', {'material': material}, ROOT/'create-submission.private.json')
        after = documents(); save_json(ROOT/'documents-after.private.json', after)
        new = [r for r in after if r['name'] == NAME and r['url'] == uploaded['url']
               and r['gateType'] == 'DOCUMENT' and r['documentKind'] == 'DOCUMENT' and r['projectId'] == 148]
        if len(new) != 1: raise ValueError('Brief registration readback failed')
        result = {'asset_id': new[0]['assetId'], 'project_id': 148, 'name': NAME, **uploaded}
        save_json(ROOT/'publication-result.json', result); print('REVIEW_BRIEF', result['asset_id'], uploaded['url'], flush=True)


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--mcp-config', type=Path, required=True); a = p.parse_args()
    if sys.platform != 'win32': p.error('Business execution is Windows-only')
    publish(a.mcp_config)
