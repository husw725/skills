"""Build and publish a version-bound bilingual review document for the revised season."""
import argparse,sys,time
from pathlib import Path
from zipfile import ZipFile
from translate import read_json,save_json,fingerprint,load_segments,srt_time
from dub_episode import digest
from process_lock import process_lock
from publish_season_candidates import candidates,PUBLICATION
from rerun_carmilla_season import sources,validate_reviewed_plan
from publish_review import materials,register,check_matches
from mflix_client import Mflix

ROOT=Path('output/review-publication/148/Carmilla/revised-comparison-20261020')


def inventory():
    available=candidates()
    if set(available)!=set(range(1,33)):raise ValueError('All 32 actual revised candidates required')
    episodes=[]
    for n,folder in sorted(available.items()):
        plan=read_json(folder/'plan.json');validate_reviewed_plan(folder,plan)
        reviewed=read_json(folder/'reviewed.json');report=read_json(folder/'render-report.json')
        source_path=sources(n)[1];original=read_json(source_path)['segments']
        subtitle_path=folder/f'Carmilla_EP{n:02}_pt-BR_revised.srt';subtitles=load_segments(subtitle_path)
        receipt=read_json(PUBLICATION/f'episode-{n:02}'/report['output_sha256']/'publication-result.json')
        from resume_season_publication import retained_delivery
        retained_delivery(n,folder,receipt,receipt,PUBLICATION/f'episode-{n:02}'/report['output_sha256']/'candidate.json')
        if (plan['episode']!=n or plan['project_id']!=59 or plan['blocking_issues']
                or digest(folder/'reviewed.json')!=plan['reviewed_response_sha256']
                or digest(source_path)!=plan['input_sha256'][str(source_path)]
                or receipt['episode']!=n or receipt['video_sha256']!=report['output_sha256']
                or digest(subtitle_path)!=receipt['subtitle_sha256'] or digest(subtitle_path)!=plan['translation_srt_sha256']
                or not receipt.get('prior_versions_retained') or report['new_clone_count']!=0
                or not report['video_payload_identical'] or report['source_audio_in_output']):
            raise ValueError('Latest source/translation/video/subtitle binding differs')
        if [c['id'] for c in reviewed['cues']]!=[r['id'] for r in original] or len(subtitles)!=len(original):
            raise ValueError('Per-cue source/translation coverage differs')
        by={r['id']:r for r in original};notes={}
        for unit in plan['units']:
            members=[by[c] for c in unit['cue_ids']]
            if unit['source_text']!=' '.join(r['text'] for r in members) or unit['source_windows']!=[{k:r[k] for k in ('id','start','end')} for r in members]:
                raise ValueError('Original subtitle text/window changed')
            for cue,evidence in unit.get('editorial_source_reconstructions',{}).items():
                kind=evidence['source_basis'];reading=evidence['editorial_reading']
                label='按 final 编辑定稿（非原声逐字确认）' if kind=='user_authorized_final_storyboard_resolution' else '编辑参考（未人工听审）'
                notes[cue]=label+'：'+reading
            for cue,evidence in unit.get('source_audio_readings',{}).items():
                if evidence.get('status')=='transcription_corrected':notes[cue]='语音识别参考（未人工听审）：'+evidence['corroborated_audio_reading']
        rows=[]
        for source,translated,subtitle in zip(original,reviewed['cues'],subtitles):
            if (abs(source['start']-subtitle['start'])>.00051 or abs(source['end']-subtitle['end'])>.00051
                    or ' '.join(translated['translation'].split())!=subtitle['text']):
                raise ValueError('Document must match delivered SRT text and original timecodes')
            rows.append({'id':source['id'],'start':source['start'],'end':source['end'],'original':source['text'],
                         'translation':subtitle['text'],'editorial_note':notes.get(source['id'],'')})
        episodes.append({'episode':n,'rows':rows,'version':receipt['version'],'video_asset_id':receipt['asset_id'],
            'subtitle_asset_id':receipt['subtitle_asset_id'],'video_sha256':receipt['video_sha256'],'subtitle_sha256':receipt['subtitle_sha256'],
            'source_sha256':digest(source_path),'plan_sha256':digest(folder/'plan.json'),'reviewed_sha256':digest(folder/'reviewed.json')})
    return episodes


def build(episodes):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ROOT.mkdir(parents=True,exist_ok=True);manifest_path=ROOT/'manifest.private.json';path=ROOT/'Carmilla_EN_pt-BR_Revised_Subtitle_Comparison.docx'
    signature=fingerprint(episodes)
    if manifest_path.exists():
        old=read_json(manifest_path)
        if old['input_fingerprint']!=signature or digest(path)!=old['document_sha256']:raise ValueError('Existing comparison checkpoint changed; preserve previous output')
        return path,old
    doc=Document();section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(1.6);section.left_margin=section.right_margin=Cm(1.5)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    normal.paragraph_format.space_after=Pt(3)
    doc.add_heading('Carmilla｜新版英葡字幕对照（32集）',0)
    doc.add_paragraph('原字幕（英文） ↔ 最新已交付巴西葡语字幕（pt-BR）')
    doc.add_paragraph('对应 Review「PT Transalte」（148）各集最新视频版本。右栏逐条采用该版本已交付字幕，不是旧版译稿。左栏保留原字幕来源与原时间码；个别按说话人拆分的字幕按制作分句列出。')
    doc.add_paragraph('标注“编辑参考”的条目保留原字幕，同时说明修订依据。第8、29集相关歧义按用户指定的 final 分镜定稿，不代表原声逐字听审通过。字幕时间码是显示/制作窗口，不是词级语音对齐。')
    doc.add_paragraph('巴葡发音、口音、情绪表演、音色衔接及混音仍待专业听审。请结合对应视频时间点留言。')
    for episode in episodes:
        n=episode['episode'];heading=doc.add_heading(f'EP.{n:02}',1);heading.paragraph_format.page_break_before=True
        doc.add_paragraph(f'对应视频：EP.{n:02} · v{episode["version"]}；本节字幕与该版本一致。')
        table=doc.add_table(rows=1,cols=2);table.style='Table Grid';table.autofit=False
        for column in table.columns:column.width=Cm(9)
        for cell,title in zip(table.rows[0].cells,['原字幕（英文）','新版译文（巴西葡语）']):cell.text=title
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for index,row in enumerate(episode['rows'],1):
            cells=table.add_row().cells;label=f'#{index}  {srt_time(row["start"])} → {srt_time(row["end"])}'
            for cell in cells:
                cell.width=Cm(9);run=cell.paragraphs[0].add_run(label);run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('637083')
            cells[0].add_paragraph(row['original']);cells[1].add_paragraph(row['translation'])
            if row['editorial_note']:
                p=cells[0].add_paragraph(row['editorial_note'])
                for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('637083')
    doc.save(path)
    reopened=Document(path)
    if len(reopened.tables)!=32:raise ValueError('Persisted document lacks an episode')
    for table,episode in zip(reopened.tables,episodes):
        if len(table.rows)!=len(episode['rows'])+1 or len(table.columns)!=2:raise ValueError('Persisted table coverage differs')
        for cells,row in zip([r.cells for r in table.rows[1:]],episode['rows']):
            if cells[0].paragraphs[1].text!=row['original'] or cells[1].paragraphs[1].text!=row['translation']:raise ValueError('Persisted left/right document text differs')
            if row['editorial_note'] and row['editorial_note'] not in cells[0].text:raise ValueError('Editorial provenance annotation missing')
    with ZipFile(path) as z:
        if z.testzip():raise ValueError('Invalid DOCX archive')
    manifest={'input_fingerprint':signature,'document_sha256':digest(path),'episodes':episodes,
              'episode_count':32,'cue_count':sum(len(e['rows']) for e in episodes),'native_listening_passed':False}
    save_json(manifest_path,manifest);return path,manifest


def publish(path,manifest,config):
    import requests
    with process_lock(ROOT/'publication.lock',timeout=0):
        if digest(path)!=manifest['document_sha256']:raise ValueError('Document changed after build')
        client=Mflix(config);ledger_path=ROOT/'review-ledger.private.json';ledger=read_json(ledger_path) if ledger_path.exists() else {'items':{}}
        from notify_season_complete import verify_delivery
        published={str(e['episode']):read_json(PUBLICATION/f'episode-{e["episode"]:02}'/e['video_sha256']/'publication-result.json') for e in manifest['episodes']}
        for e in manifest['episodes']:
            receipt=published[str(e['episode'])]
            for key in ('video_sha256','subtitle_sha256','version','subtitle_asset_id'):
                if receipt[key]!=e[key]:raise ValueError('Document delivery binding changed')
            if receipt['asset_id']!=e['video_asset_id']:raise ValueError('Document video binding changed')
        verify_delivery({'all_32_published':True,'published':published},materials(client,148))
        sha=manifest['document_sha256'];upload=ROOT/'upload.private.json'
        if upload.exists():
            cached=read_json(upload)
            if cached['sha256']!=sha:raise ValueError('Prior upload belongs to another document')
        else:
            signed=client.call('createFilePresignedUrl',{'originalFileName':path.name})
            if not isinstance(signed,str) or not signed.startswith('https://'):raise ValueError('Expected Review HTTPS upload URL')
            cached={'sha256':sha,'upload_url':signed,'url':signed.split('?')[0],'status':'uploading'};save_json(upload,cached)
        if cached['status']!='uploaded':
            import hashlib
            body=path.read_bytes()
            if hashlib.sha256(body).hexdigest()!=sha:raise ValueError('Document changed before upload')
            try:
                response=client.http.put(cached['upload_url'],data=body,timeout=90);response.raise_for_status()
            except requests.RequestException:raise RuntimeError('Document upload failed; signed object checkpoint retained') from None
            cached['status']='uploaded';save_json(upload,cached)
        spec={'projectId':148,'gateType':'DOCUMENT','documentKind':'DOCUMENT',
              'name':f'Carmilla | 新版英葡字幕对照（32集） | {sha[:12]}.docx','url':cached['url'],
              'note':f'Latest revised season: 32 episodes, {manifest["cue_count"]} aligned subtitle rows. Original English on left; current delivered pt-BR on right. User-final editorial decisions annotated; native listening pending. SHA256: {sha}'}
        asset=register(client,spec,ledger,ledger_path);row=check_matches(materials(client,148),spec)
        if row['assetId']!=asset or row.get('documentKind')!='DOCUMENT' or sha not in row.get('note',''):raise ValueError('Review document readback differs')
        try:
            response=requests.get(cached['url'],timeout=60);response.raise_for_status();body=response.content
        except requests.RequestException:raise RuntimeError('Document download verification failed; registration receipts retained') from None
        import hashlib
        if hashlib.sha256(body).hexdigest()!=sha:raise ValueError('Downloaded document differs')
        save_json(ROOT/'publication-result.json',{'project_id':148,'asset_id':asset,'name':spec['name'],'url':cached['url'],
            'sha256':sha,'episodes':32,'cue_count':manifest['cue_count'],'verified_at':time.time()})
        print('NEW_COMPARISON_DOCUMENT_VERIFIED',asset,manifest['cue_count'],flush=True)


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--publish',action='store_true');p.add_argument('--mcp-config',type=Path);a=p.parse_args()
    if sys.platform!='win32':p.error('Business processing runs on Windows')
    if a.publish and (not a.mcp_config or not a.mcp_config.is_file()):p.error('Existing Review configuration required')
    path,manifest=build(inventory());print('NEW_COMPARISON_BUILT',manifest['episode_count'],manifest['cue_count'],manifest['document_sha256'],flush=True)
    if a.publish:publish(path,manifest,a.mcp_config)


if __name__=='__main__':main()
